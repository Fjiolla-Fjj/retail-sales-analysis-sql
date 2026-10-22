"""Reproduce the full analysis: load data -> SQLite -> run SQL files -> Excel workbook + charts.
Run from repo root:  python scripts/analysis.py   (needs pandas, openpyxl, matplotlib)"""
import sqlite3, re, json
from pathlib import Path
import pandas as pd
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

ROOT = Path(__file__).resolve().parents[1]
con = sqlite3.connect(":memory:")
df = pd.read_csv(ROOT / "data/online_retail.csv.gz", dtype={"InvoiceNo": str, "StockCode": str, "CustomerID": str})
df.to_sql("raw_sales", con, index=False)
con.execute("CREATE INDEX ix_raw ON raw_sales(CustomerID, StockCode, Quantity)")  # speeds up the cancellation match

def run(fname):
    """Split a .sql file on '-- Qn' headers and run each block; return {title: DataFrame}."""
    text = (ROOT / "sql" / fname).read_text()
    out = {}
    for blk in re.split(r"\n(?=-- Q\d)", text):
        title = blk.splitlines()[0].lstrip("- ").strip()
        stmts = [s.strip() for s in blk.split(";") if re.sub(r"--.*", "", s).strip()]
        for s in stmts:
            cur = con.execute(s)
            if cur.description:
                out[title] = pd.DataFrame(cur.fetchall(), columns=[d[0] for d in cur.description])
    return out

R = {}
for f in ["01_data_cleaning_and_validation.sql", "02_kpis.sql", "03_customer_product_analysis.sql"]:
    res = run(f); R.update(res)
    print(f"\n===== {f} =====")
    for t, d in res.items():
        print(f"\n-- {t}\n{d.to_string(index=False)}")

def get(prefix): return next(v for k, v in R.items() if k.startswith(prefix))
kpi = get("Q1 Headline"); monthly = get("Q2 Monthly"); top_p = get("Q1 Top 10 products")
top_c = get("Q2 Top 10 customers"); country = get("Q3 Country"); rfm = get("Q4 RFM")

# ---- charts ----
img = ROOT / "images"; img.mkdir(exist_ok=True)
plt.rcParams.update({"figure.dpi": 120, "axes.spines.top": False, "axes.spines.right": False})
fig, ax = plt.subplots(figsize=(10, 4.5))
ax.plot(monthly.month, monthly.revenue / 1000, marker="o", color="#1f6feb")
ax.set_title("Monthly Revenue (GBP '000) - Dec 2011 is a partial month (to 9 Dec)"); ax.set_ylabel("Revenue (£k)")
plt.xticks(rotation=45); plt.tight_layout(); plt.savefig(img / "monthly_revenue.png"); plt.close()

fig, ax = plt.subplots(figsize=(10, 5))
tp = top_p.iloc[::-1]; ax.barh(tp["product"].str.title().str[:35], tp.revenue / 1000, color="#2da44e")
ax.set_title("Top 10 Products by Revenue (£k)"); plt.tight_layout(); plt.savefig(img / "top_products.png"); plt.close()

fig, ax = plt.subplots(figsize=(10, 4.5))
cc = country.iloc[::-1]; ax.barh(cc.Country, cc.revenue / 1000, color="#bf8700")
ax.set_xscale("log"); ax.set_title("Revenue by Country - Top 10 (£k, log scale)"); plt.tight_layout()
plt.savefig(img / "revenue_by_country.png"); plt.close()

fig, ax = plt.subplots(figsize=(10, 4.5))
ax.bar(rfm.segment, rfm.revenue_share_pct, color="#8250df")
for i, (s, c) in enumerate(zip(rfm.revenue_share_pct, rfm.customers)): ax.text(i, s + 1, f"{s}%\n{c} cust.", ha="center", fontsize=9)
ax.set_title("RFM Segments - Share of Revenue"); ax.set_ylabel("% of revenue"); ax.set_ylim(0, rfm.revenue_share_pct.max() + 12)
plt.tight_layout(); plt.savefig(img / "rfm_segments.png"); plt.close()

# ---- Excel workbook ----
wb = Workbook(); hdr = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1F4E78")
def sheet(name, d, ws=None):
    ws = ws or wb.create_sheet(name)
    for r in dataframe_to_rows(d, index=False, header=True): ws.append(r)
    for c in ws[1]: c.font, c.fill = hdr, fill
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = max(12, min(45, max(len(str(c.value)) for c in col) + 2))
    return ws
ws = wb.active; ws.title = "KPIs"
sheet(None, kpi.T.reset_index().rename(columns={"index": "KPI", 0: "Value"}), ws)
ws = sheet("Monthly Trend", monthly)
ch = LineChart(); ch.title = "Monthly Revenue (GBP)"; ch.height, ch.width = 8, 18
ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(monthly) + 1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(monthly) + 1)); ws.add_chart(ch, "H2")
ws = sheet("Top 10 Products", top_p)
bc = BarChart(); bc.type = "bar"; bc.title = "Top 10 Products by Revenue"; bc.height, bc.width = 9, 18
bc.add_data(Reference(ws, min_col=4, min_row=1, max_row=11), titles_from_data=True)
bc.set_categories(Reference(ws, min_col=2, min_row=2, max_row=11)); ws.add_chart(bc, "F2")
sheet("Top 10 Customers", top_c); sheet("Top 10 Countries", country); sheet("RFM Segments", rfm)
(ROOT / "excel").mkdir(exist_ok=True); wb.save(ROOT / "excel/retail_sales_summary.xlsx")
print("\nSaved excel/retail_sales_summary.xlsx and images/*.png")
